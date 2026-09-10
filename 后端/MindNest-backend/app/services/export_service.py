from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.question import Question
from app.models.recognition_task import RecognitionTask
from app.models.user import User
from app.models.wrong_record import WrongRecord
from app.schemas.export import ExportRequest
from app.services.question_service import list_questions
from app.services.serializers import question_out
from app.utils.time_utils import format_datetime, now


def _questions(db: Session, current_user: User, payload: ExportRequest) -> list[dict]:
    export_range = (payload.range or "current").lower()
    if payload.taskId:
        statement = select(Question).where(Question.task_id == payload.taskId).order_by(Question.created_at.asc())
        if current_user.role != "admin":
            statement = statement.where(Question.user_id == current_user.id)
        return [question_out(item) for item in db.scalars(statement).all()]

    if export_range in {"wrong", "mastery", "batch", "knowledge"}:
        statement = select(Question)
        if export_range in {"wrong", "mastery"}:
            statement = statement.join(WrongRecord, WrongRecord.question_id == Question.id)
            if payload.masteryStatus:
                statement = statement.where(WrongRecord.mastery_status == payload.masteryStatus)
            if current_user.role != "admin":
                statement = statement.where(WrongRecord.user_id == current_user.id)
        if export_range == "batch" and payload.batchId:
            statement = statement.join(RecognitionTask, RecognitionTask.id == Question.task_id).where(RecognitionTask.batch_id == payload.batchId)
        if export_range == "knowledge" and payload.knowledgePoint:
            statement = statement.where(Question.knowledge_points.like(f"%{payload.knowledgePoint}%"))
        if payload.subject:
            statement = statement.where(Question.subject == payload.subject)
        if payload.questionType:
            statement = statement.where(Question.question_type == payload.questionType)
        if payload.difficulty:
            statement = statement.where(Question.difficulty == payload.difficulty)
        if current_user.role != "admin" and export_range not in {"wrong", "mastery"}:
            statement = statement.where(Question.user_id == current_user.id)
        statement = statement.order_by(Question.created_at.asc()).distinct()
        return [question_out(item) for item in db.scalars(statement).all()]

    rows, _, _, _ = list_questions(
        db,
        current_user=current_user,
        keyword=payload.keyword,
        subject=payload.subject,
        question_type=payload.questionType,
        difficulty=payload.difficulty,
        knowledge_point=payload.knowledgePoint,
        start_time=payload.startTime,
        end_time=payload.endTime,
        page=1,
        page_size=1000,
        all_users=current_user.role == "admin",
    )
    return rows


def _as_lines(value) -> list[str]:
    if not value:
        return []
    if isinstance(value, list):
        lines = []
        for item in value:
            if isinstance(item, dict):
                text = item.get("text") or item.get("content") or item.get("latex") or item.get("label") or str(item)
            else:
                text = str(item)
            if text:
                lines.append(str(text))
        return lines
    return [str(value)]


def _markdown(rows: list[dict], payload: ExportRequest) -> str:
    lines = [
        "# 智学空间(MindNest) 题目整理导出",
        "",
        f"> 导出时间：{format_datetime(now())}",
        f"> 范围：{payload.range}",
        f"> 格式：{payload.format.upper()}",
        "",
    ]
    for index, item in enumerate(rows, start=1):
        vlm = item.get("vlmDetail") or {}
        options = item.get("options") or []
        formulas = item.get("formulas") or vlm.get("formulas") or []
        lines += [
            f"## 题目 {index} · {item.get('title', '')}",
            f"- 科目：{item.get('subject', '')}",
            f"- 题型：{item.get('questionType', '')}",
            f"- 难度：{item.get('difficulty', '')}",
            f"- 置信度：{item.get('confidence', '')}",
            f"- 知识点：{' / '.join(item.get('knowledgePoints') or [])}",
            "",
            "### 题干",
            item.get("content", ""),
            "",
        ]
        if options:
            lines.append("### 选项")
            for option in options:
                if isinstance(option, dict):
                    label = option.get("label") or option.get("key") or ""
                    text = option.get("text") or option.get("content") or ""
                    lines.append(f"- {label} {text}".strip())
                else:
                    lines.append(f"- {option}")
            lines.append("")
        if formulas:
            lines.append("### 公式")
            for formula in formulas:
                if isinstance(formula, dict):
                    lines.append(f"- {formula.get('latex') or formula.get('text') or formula}")
                else:
                    lines.append(f"- {formula}")
            lines.append("")
        lines += [
            "### OCR 识别文本",
            item.get("ocrText") or "暂无",
            "",
            "### VLM 解析",
            f"- 答案：{vlm.get('answer') or item.get('answer') or '暂无'}",
            f"- 解析摘要：{vlm.get('analysisSummary') or item.get('analysisSummary') or '暂无'}",
        ]
        detail = vlm.get("detailedAnalysis") or item.get("detailedAnalysis")
        if detail:
            lines += ["", "#### 详细解析", str(detail)]
        for title, key in (
            ("解题步骤", "solutionSteps"),
            ("常见错误", "commonMistakes"),
            ("错误原因", "errorCauseTags"),
            ("复习计划", "reviewPlan"),
            ("相似练习建议", "similarPracticeSuggestions"),
        ):
            values = _as_lines(vlm.get(key) or item.get(key))
            if values:
                lines += ["", f"#### {title}"]
                lines += [f"- {value}" for value in values]
        lines.append("")
    return "\n".join(lines)


def _word(markdown: str) -> bytes:
    doc = Document()
    for line in markdown.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("> "):
            doc.add_paragraph(line[2:])
        elif line:
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _register_font() -> str:
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    ]
    for path in candidates:
        if path.exists():
            pdfmetrics.registerFont(TTFont("MindNestCN", str(path)))
            return "MindNestCN"
    return "Helvetica"


def _pdf(markdown: str) -> bytes:
    font_name = _register_font()
    buffer = BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    _, height = A4
    x, y = 48, height - 48
    page.setFont(font_name, 11)
    for raw_line in markdown.splitlines():
        line = raw_line.replace("#", "").strip()
        if not line:
            y -= 10
            continue
        for start in range(0, len(line), 56):
            if y < 48:
                page.showPage()
                page.setFont(font_name, 11)
                y = height - 48
            page.drawString(x, y, line[start:start + 56])
            y -= 18
    page.save()
    return buffer.getvalue()


def build_export(db: Session, current_user: User, payload: ExportRequest) -> tuple[bytes, str, str]:
    rows = _questions(db, current_user, payload)
    markdown = _markdown(rows, payload)
    fmt = (payload.format or "markdown").lower()
    if fmt in {"word", "docx"}:
        return _word(markdown), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "mindnest-export.docx"
    if fmt == "pdf":
        return _pdf(markdown), "application/pdf", "mindnest-export.pdf"
    return markdown.encode("utf-8"), "text/markdown; charset=utf-8", "mindnest-export.md"
